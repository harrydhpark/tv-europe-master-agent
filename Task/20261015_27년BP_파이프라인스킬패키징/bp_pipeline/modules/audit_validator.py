# -*- coding: utf-8 -*-
"""
audit_validator.py:
Performs independent full-scale verification of 2027 Business Plan XLSB files.
Directly reads all European rows (7,361 rows) using pyxlsb.
"""
import os
import re
from collections import defaultdict
import openpyxl
from pyxlsb import open_workbook

class AuditValidator:
    def __init__(self, xlsb_path, plan_path, q1_oled_targets=None):
        self.xlsb_path = xlsb_path
        self.plan_path = plan_path
        self.q1_oled_targets = q1_oled_targets or {
            'AG': 10754, 'BN': 27113, 'CK': 11319, 'DG': 49594,
            'Swiss': 9714, 'ES': 27881, 'FS': 27108, 'HS': 3185,
            'IS': 31324, 'LA': 6936, 'MK': 18394, 'PL': 14412,
            'PT': 6274, 'RO': 6572, 'SW': 23139, 'UK': 63177
        }
        self.load_plan_targets()

    def load_plan_targets(self):
        wb_plan = openpyxl.load_workbook(self.plan_path, data_only=True)
        ws_p = wb_plan["4. 법인별 수량_시나리오별"]
        self.subs = [ws_p.cell(r, 7).value for r in range(6, 22)]

        self.strat_targets = {}
        self.oled_97_targets = {}
        self.oled_77up_targets = {}
        self.uhd_giant_targets = {}

        for r in range(6, 22):
            sub = ws_p.cell(r, 7).value
            oled = ws_p.cell(r, 16).value or 0
            w = ws_p.cell(r, 21).value or 0
            g = ws_p.cell(r, 26).value or 0
            cb = oled - w - g
            t_97 = ws_p.cell(r, 31).value or 0
            t_77up = ws_p.cell(r, 36).value or 0
            micro = ws_p.cell(r, 51).value or 0
            mini = ws_p.cell(r, 56).value or 0
            qned = ws_p.cell(r, 61).value or 0
            uhd = ws_p.cell(r, 66).value or 0
            giant = ws_p.cell(r, 71).value or 0
            fhd = ws_p.cell(r, 81).value or 0

            self.strat_targets[(sub, 'W시리즈')] = int(round(w))
            self.strat_targets[(sub, 'G시리즈')] = int(round(g))
            self.strat_targets[(sub, 'CB시리즈')] = int(round(cb))
            self.strat_targets[(sub, 'Micro RGB')] = int(round(micro))
            self.strat_targets[(sub, 'Mini RGB')] = int(round(mini))
            self.strat_targets[(sub, 'QNED')] = int(round(qned))
            self.strat_targets[(sub, '일반UHD')] = int(round(uhd))
            self.strat_targets[(sub, 'FHD')] = int(round(fhd))
            self.strat_targets[(sub, 'OLED')] = int(round(oled))
            self.strat_targets[(sub, 'UHD')] = int(round(micro + mini + qned + uhd))

            self.oled_97_targets[sub] = int(round(t_97))
            self.oled_77up_targets[sub] = int(round(t_77up))
            self.uhd_giant_targets[sub] = int(round(giant))

        self.target_financials = {}
        for r in range(67, 83):
            s = ws_p.cell(r, 7).value
            gross_m = ws_p.cell(r, 8).value or 0
            net_m = ws_p.cell(r, 9).value or 0
            sd_r = ws_p.cell(r, 10).value or 0
            self.target_financials[s] = {
                'gross': gross_m * 1e6,
                'net': net_m * 1e6,
                'sd_rate': sd_r
            }
        wb_plan.close()

    @staticmethod
    def safe_inch(val):
        if val is None: return 0
        s = str(val).strip()
        try: return int(round(float(s)))
        except:
            m = re.search(r'\d+', s)
            return int(m.group(0)) if m else 0

    @staticmethod
    def get_segment(blu, prod, tool):
        if blu == 'OLED':
            if tool.startswith('W') or tool.startswith('M') or tool.startswith('Z'): return 'W시리즈'
            elif tool.startswith('G'): return 'G시리즈'
            else: return 'CB시리즈'
        elif blu == 'UHD':
            if prod == 'MRGB':
                if tool.startswith('MRGB9'): return 'Micro RGB'
                else: return 'Mini RGB'
            elif prod == 'QNED': return 'QNED'
            else: return '일반UHD'
        elif blu == 'FHD': return 'FHD'
        return 'OTHER'

    def run_audit(self):
        rows = []
        with open_workbook(self.xlsb_path) as wb:
            with wb.get_sheet('②모델별RAW') as sheet:
                for idx, r in enumerate(sheet.rows()):
                    if idx < 15: continue
                    vals = [c.v for c in r]
                    if len(vals) > 23 and str(vals[9]).strip() == '02.유럽':
                        sub_code = str(vals[11]).strip() if vals[11] else ''
                        ctr_code = str(vals[15]).strip() if vals[15] else ''
                        sub_key = 'Swiss' if sub_code == 'LGEDG' and ctr_code == 'Swiss' else sub_code.replace('LGE', '')
                        blu = str(vals[16]).strip().upper() if vals[16] else ''
                        prod = str(vals[17]).strip().upper() if vals[17] else ''
                        inch = self.safe_inch(vals[18])
                        model = str(vals[19]).strip() if vals[19] else ''
                        yr = '26' if '26' in str(vals[22]) else '27'
                        dist_no = str(vals[6]).strip() if vals[6] else ''
                        tool = str(vals[23]).strip().upper() if vals[23] else ''
                        seg = self.get_segment(blu, prod, tool)

                        def safe_f(v):
                            try: return float(v or 0)
                            except: return 0.0

                        m_qty = [int(round(safe_f(vals[40+m]))) for m in range(12)]
                        fx_rates = [safe_f(vals[28+m]) or 1.0 for m in range(12)]
                        base_gp = [safe_f(vals[57+m]) for m in range(12)]
                        base_sd = [safe_f(vals[69+m]) for m in range(12)]

                        rows.append({
                            'row': idx + 1,
                            'sub_key': sub_key,
                            'blu': blu,
                            'prod': prod,
                            'inch': inch,
                            'tool': tool,
                            'model': model,
                            'yr': yr,
                            'dist_no': dist_no,
                            'seg': seg,
                            'm_qty': m_qty,
                            'fx_rates': fx_rates,
                            'base_gp': base_gp,
                            'base_sd': base_sd
                        })

        tot_vol = sum(sum(r['m_qty']) for r in rows)
        
        q1_oled_act = defaultdict(int)
        for s in self.subs:
            q1_oled_act[s] = sum(sum(r['m_qty'][:3]) for r in rows if r['sub_key'] == s and r['blu'] == 'OLED')

        uhd_giant_act = defaultdict(int)
        for s in self.subs:
            uhd_giant_act[s] = sum(sum(r['m_qty']) for r in rows if r['sub_key'] == s and r['blu'] == 'UHD' and r['inch'] >= 98)

        tot_act_97 = 0
        tot_act_77 = 0
        for s in self.subs:
            tot_act_97 += sum(sum(r['m_qty']) for r in rows if r['sub_key'] == s and r['blu'] == 'OLED' and r['inch'] == 97)
            tot_act_77 += sum(sum(r['m_qty']) for r in rows if r['sub_key'] == s and r['blu'] == 'OLED' and r['inch'] >= 77)

        strat_mismatches = 0
        for s in self.subs:
            for seg in ['W시리즈', 'G시리즈', 'CB시리즈', 'Micro RGB', 'Mini RGB', 'QNED', '일반UHD', 'FHD']:
                act = sum(sum(r['m_qty']) for r in rows if r['sub_key'] == s and r['seg'] == seg)
                tgt = self.strat_targets.get((s, seg), 0)
                if act != tgt:
                    strat_mismatches += 1

        fin_gross, fin_net = 0.0, 0.0
        for r in rows:
            for m in range(12):
                q = r['m_qty'][m]
                gp = r['base_gp'][m]
                fx = r['fx_rates'][m]
                sd = r['base_sd'][m]
                gsp = gp / fx if fx else gp
                g = q * gsp
                n = g * (1.0 - sd)
                fin_gross += g
                fin_net += n
        fin_sd = (fin_gross - fin_net) / fin_gross * 100 if fin_gross > 0 else 0.0

        v_27_m14 = sum(sum(r['m_qty'][:4]) for r in rows if r['yr'] == '27')
        v_26_m512 = sum(sum(r['m_qty'][4:]) for r in rows if r['yr'] == '26' and r['seg'] != 'FHD' and '98NU900' not in r['model'] and '115QNED' not in r['model'])
        missing_gp = sum(1 for r in rows for m in range(12) if r['m_qty'][m] > 0 and r['base_gp'][m] <= 0)
        missing_sd = sum(1 for r in rows for m in range(12) if r['m_qty'][m] > 0 and r['base_sd'][m] <= 0)

        q1_diff = {s: q1_oled_act[s] - self.q1_oled_targets.get(s, 0) for s in self.subs}
        giant_diff = {s: uhd_giant_act[s] - self.uhd_giant_targets.get(s, 0) for s in self.subs}

        report = {
            'total_volume': tot_vol,
            'total_volume_diff': tot_vol - 5264650,
            'q1_oled_total': sum(q1_oled_act.values()),
            'q1_oled_total_diff': sum(q1_oled_act.values()) - sum(self.q1_oled_targets.values()),
            'q1_oled_diff_by_sub': q1_diff,
            'uhd_giant_total': sum(uhd_giant_act.values()),
            'uhd_giant_total_diff': sum(uhd_giant_act.values()) - sum(self.uhd_giant_targets.values()),
            'uhd_giant_diff_by_sub': giant_diff,
            'oled_97_total': tot_act_97,
            'oled_97_diff': tot_act_97 - 2501,
            'oled_77up_total': tot_act_77,
            'oled_77up_diff': tot_act_77 - 223300,
            'strat_mismatches': strat_mismatches,
            'total_gross_m': fin_gross / 1e6,
            'total_net_m': fin_net / 1e6,
            'net_diff_m': (fin_net / 1e6) - 2545.20,
            'total_sd_rate': fin_sd,
            'clean_cut_violations_27yr_m1_m4': v_27_m14,
            'clean_cut_violations_26yr_m5_m12': v_26_m512,
            'missing_gp_count': missing_gp,
            'missing_sd_count': missing_sd,
            'is_perfect': (
                abs(tot_vol - 5264650) == 0 and
                all(d == 0 for d in q1_diff.values()) and
                all(d == 0 for d in giant_diff.values()) and
                strat_mismatches == 0 and
                abs((fin_net / 1e6) - 2545.20) < 0.01 and
                abs(fin_sd - 28.0000) < 0.01 and
                v_27_m14 == 0 and v_26_m512 == 0 and
                missing_gp == 0 and missing_sd == 0
            )
        }
        return report
