import os
import sys
import json
import sqlite3
import openpyxl

ROOT_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
DATA_DIR = os.path.join(ROOT_DIR, "data")

def patch_raw_mediamarkt_lg():
    raw_path = os.path.join(DATA_DIR, "raw_mediamarkt_lg.json")
    with open(raw_path, "r", encoding="utf-8") as f:
        items = json.load(f)

    updated = False
    for it in items:
        if it.get("model_code") == "OLED65C67LA" or "2291289" in it.get("link", ""):
            it["price"] = 1599.0
            it["promo"] = "Aus unserer Werbung | 25% LG Sound Suite"
            updated = True
            print("  [RAW LG] Updated OLED65C67LA to 1599.0 CHF")

    with open(raw_path, "w", encoding="utf-8") as f:
        json.dump(items, f, indent=2, ensure_ascii=False)
    print(f"[RAW LG] Saved {raw_path} (updated={updated})")

def patch_price_tracker_workbook():
    pt_path = os.path.join(DATA_DIR, "price tracker_swiss_2026 1006_v1.xlsx")
    if not os.path.exists(pt_path):
        print(f"[WARN] {pt_path} not found.")
        return

    wb = openpyxl.load_workbook(pt_path)
    sheet = wb["MediaMarkt_LG_Full"]

    for r in range(2, sheet.max_row + 1):
        code = str(sheet.cell(row=r, column=5).value or "")
        if "OLED65C67LA" in code:
            sheet.cell(row=r, column=6).value = 1599.0
            sheet.cell(row=r, column=10).value = "Aus unserer Werbung"
            print(f"  [PT WORKBOOK] Updated row {r} (OLED65C67LA) price to 1599.0 CHF")

    wb.save(pt_path)
    wb.close()
    print(f"[PT WORKBOOK] Successfully updated {pt_path}")

def patch_swiss_ata_comparison():
    ata_path = os.path.join(DATA_DIR, "Swiss_ATA_Comparison_2026_1006.xlsx")
    if not os.path.exists(ata_path):
        print(f"[WARN] {ata_path} not found.")
        return

    wb = openpyxl.load_workbook(ata_path)
    sheet = wb["Swiss_2026"]

    for r in range(4, sheet.max_row + 1):
        val_b = str(sheet.cell(row=r, column=2).value or "").strip()
        val_c = str(sheet.cell(row=r, column=3).value or "").strip()
        if "65C6" in val_b and "OLED65C67LA" in val_c:
            sheet.cell(row=r, column=5).value = 1599.0
            print(f"  [ATA WORKBOOK] Updated row {r} Col E to 1599.0 CHF")

    wb.save(ata_path)
    wb.close()
    print(f"[ATA WORKBOOK] Successfully updated {ata_path}")

def patch_price_history_db():
    db_path = os.path.join(DATA_DIR, "price_history.db")
    conn = sqlite3.connect(db_path)
    cur = conn.cursor()
    cur.execute("""
        UPDATE daily_prices 
        SET selling_price = 1599.0, promo_text = 'Aus unserer Werbung | 25% LG Sound Suite'
        WHERE country = 'CH' AND retailer = 'MediaMarkt' AND model_code = 'OLED65C67LA' AND date = '2026-10-06'
    """)
    print(f"  [DB] Updated {cur.rowcount} rows in daily_prices")
    conn.commit()
    conn.close()

if __name__ == "__main__":
    print("=== Step 1: Update raw_mediamarkt_lg.json ===")
    patch_raw_mediamarkt_lg()

    print("\n=== Step 2: Update price tracker_swiss_2026 1006_v1.xlsx ===")
    patch_price_tracker_workbook()

    print("\n=== Step 3: Update Swiss_ATA_Comparison_2026_1006.xlsx ===")
    patch_swiss_ata_comparison()

    print("\n=== Step 4: Update price_history.db ===")
    patch_price_history_db()

    print("\nAll patch operations completed successfully.")
