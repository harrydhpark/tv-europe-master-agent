import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Subsidiary pseudonym mapping (Order by WOS descending)
SUBS_MAPPING = [
    {"code": "A", "name": "A 법인", "wos": 12.62, "risk": "과다재고 (>6주)", "inv_ttl": 29906, "inv_new": 22931, "so_ttl": 9481, "so_oled_new": 1883, "oled_in_oled": 0.790, "oled_in_new": 0.508, "oled_new_in_ttl": 0.199, "wos_new": 24.73, "wos_old": 4.83, "is_risk": True},
    {"code": "B", "name": "B 법인", "wos": 8.58, "risk": "과다재고 (>6주)", "inv_ttl": 22458, "inv_new": 20956, "so_ttl": 10468, "so_oled_new": 1724, "oled_in_oled": 0.661, "oled_in_new": 0.244, "oled_new_in_ttl": 0.165, "wos_new": 11.86, "wos_old": 1.77, "is_risk": True},
    {"code": "C", "name": "C 법인", "wos": 8.46, "risk": "과다재고 (>6주)", "inv_ttl": 49241, "inv_new": 43217, "so_ttl": 23292, "so_oled_new": 5051, "oled_in_oled": 0.521, "oled_in_new": 0.441, "oled_new_in_ttl": 0.217, "wos_new": 15.11, "wos_old": 2.03, "is_risk": True},
    {"code": "D", "name": "D 법인", "wos": 8.44, "risk": "과다재고 (>6주)", "inv_ttl": 44396, "inv_new": 28501, "so_ttl": 21032, "so_oled_new": 2262, "oled_in_oled": 0.406, "oled_in_new": 0.255, "oled_new_in_ttl": 0.108, "wos_new": 12.85, "wos_old": 5.23, "is_risk": True},
    {"code": "E", "name": "E 법인", "wos": 8.01, "risk": "과다재고 (>6주)", "inv_ttl": 77793, "inv_new": 63588, "so_ttl": 38837, "so_oled_new": 11426, "oled_in_oled": 0.572, "oled_in_new": 0.441, "oled_new_in_ttl": 0.294, "wos_new": 9.82, "wos_old": 4.39, "is_risk": True},
    {"code": "F", "name": "F 법인", "wos": 7.73, "risk": "과다재고 (>6주)", "inv_ttl": 91671, "inv_new": 72583, "so_ttl": 47443, "so_oled_new": 3551, "oled_in_oled": 0.706, "oled_in_new": 0.095, "oled_new_in_ttl": 0.075, "wos_new": 7.75, "wos_old": 7.65, "is_risk": True},
    {"code": "G", "name": "G 법인", "wos": 7.35, "risk": "과다재고 (>6주)", "inv_ttl": 24183, "inv_new": 17743, "so_ttl": 13157, "so_oled_new": 2149, "oled_in_oled": 0.379, "oled_in_new": 0.559, "oled_new_in_ttl": 0.163, "wos_new": 18.45, "wos_old": 2.77, "is_risk": True},
    {"code": "H", "name": "H 법인", "wos": 7.31, "risk": "과다재고 (>6주)", "inv_ttl": 53208, "inv_new": 29204, "so_ttl": 29116, "so_oled_new": 3360, "oled_in_oled": 0.363, "oled_in_new": 0.302, "oled_new_in_ttl": 0.115, "wos_new": 10.50, "wos_old": 5.34, "is_risk": True},
    {"code": "I", "name": "I 법인", "wos": 7.24, "risk": "과다재고 (>6주)", "inv_ttl": 14469, "inv_new": 9543, "so_ttl": 7991, "so_oled_new": 486, "oled_in_oled": 0.368, "oled_in_new": 0.157, "oled_new_in_ttl": 0.061, "wos_new": 12.32, "wos_old": 4.03, "is_risk": True},
    {"code": "J", "name": "J 법인", "wos": 7.03, "risk": "과다재고 (>6주)", "inv_ttl": 55383, "inv_new": 31787, "so_ttl": 31501, "so_oled_new": 3177, "oled_in_oled": 0.705, "oled_in_new": 0.224, "oled_new_in_ttl": 0.101, "wos_new": 8.96, "wos_old": 5.45, "is_risk": True},
    {"code": "K", "name": "K 법인", "wos": 6.39, "risk": "과다재고 (>6주)", "inv_ttl": 120916, "inv_new": 99763, "so_ttl": 75696, "so_oled_new": 8425, "oled_in_oled": 0.655, "oled_in_new": 0.160, "oled_new_in_ttl": 0.111, "wos_new": 7.58, "wos_old": 3.67, "is_risk": True},
    {"code": "L", "name": "L 법인", "wos": 6.38, "risk": "과다재고 (>6주)", "inv_ttl": 12491, "inv_new": 5891, "so_ttl": 7830, "so_oled_new": 1116, "oled_in_oled": 0.649, "oled_in_new": 0.445, "oled_new_in_ttl": 0.143, "wos_new": 9.39, "wos_old": 4.96, "is_risk": True},
    {"code": "M", "name": "M 법인", "wos": 6.31, "risk": "과다재고 (>6주)", "inv_ttl": 7681, "inv_new": 3615, "so_ttl": 4870, "so_oled_new": 920, "oled_in_oled": 0.408, "oled_in_new": 0.450, "oled_new_in_ttl": 0.189, "wos_new": 7.07, "wos_old": 5.76, "is_risk": True},
    {"code": "N", "name": "N 법인", "wos": 5.74, "risk": "정상 (3~6주)", "inv_ttl": 26296, "inv_new": 18626, "so_ttl": 18338, "so_oled_new": 4536, "oled_in_oled": 0.490, "oled_in_new": 0.523, "oled_new_in_ttl": 0.247, "wos_new": 8.58, "wos_old": 3.18, "is_risk": False},
    {"code": "O", "name": "O 법인", "wos": 5.62, "risk": "정상 (3~6주)", "inv_ttl": 17888, "inv_new": 17087, "so_ttl": 12740, "so_oled_new": 1260, "oled_in_oled": 0.886, "oled_in_new": 0.136, "oled_new_in_ttl": 0.099, "wos_new": 7.35, "wos_old": 0.93, "is_risk": False},
    {"code": "P", "name": "P 법인", "wos": 4.23, "risk": "정상 (3~6주)", "inv_ttl": 6691, "inv_new": 2173, "so_ttl": 6328, "so_oled_new": 986, "oled_in_oled": 0.509, "oled_in_new": 0.405, "oled_new_in_ttl": 0.156, "wos_new": 3.57, "wos_old": 4.64, "is_risk": False},
    {"code": "EU", "name": "유럽 본부 (EU 총괄)", "wos": 7.31, "risk": "과다재고 (>6주)", "inv_ttl": 654669, "inv_new": 487207, "so_ttl": 358121, "so_oled_new": 52312, "oled_in_oled": 0.548, "oled_in_new": 0.256, "oled_new_in_ttl": 0.146, "wos_new": 9.54, "wos_old": 4.36, "is_risk": True, "is_eu": True}
]

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "리스크법인_WOS_OLED요약"
ws.views.sheetView[0].showGridLines = True

# Colors
NAVY = "1B365D"
DARK_NAVY = "0F2027"
LIGHT_BLUE = "F0F4F8"
RISK_RED_BG = "FDE8E8"
RISK_RED_TXT = "C81E1E"
SAFE_GREEN_BG = "DEF7EC"
SAFE_GREEN_TXT = "03543F"
ZEBRA_ROW = "F9FAFB"

font_title = Font(name="Malgun Gothic", size=14, bold=True, color="1B365D")
font_sub = Font(name="Malgun Gothic", size=9, color="555555", italic=True)
font_hdr = Font(name="Malgun Gothic", size=10, bold=True, color="FFFFFF")
font_data = Font(name="Malgun Gothic", size=9)
font_bold = Font(name="Malgun Gothic", size=9, bold=True)
font_risk = Font(name="Malgun Gothic", size=9, bold=True, color=RISK_RED_TXT)
font_safe = Font(name="Malgun Gothic", size=9, bold=True, color=SAFE_GREEN_TXT)

fill_hdr = PatternFill(start_color=NAVY, end_color=NAVY, fill_type="solid")
fill_zebra = PatternFill(start_color=ZEBRA_ROW, end_color=ZEBRA_ROW, fill_type="solid")
fill_risk = PatternFill(start_color=RISK_RED_BG, end_color=RISK_RED_BG, fill_type="solid")
fill_safe = PatternFill(start_color=SAFE_GREEN_BG, end_color=SAFE_GREEN_BG, fill_type="solid")
fill_total = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")

thin_border = Border(
    left=Side(style='thin', color='D1D5DB'),
    right=Side(style='thin', color='D1D5DB'),
    top=Side(style='thin', color='D1D5DB'),
    bottom=Side(style='thin', color='D1D5DB')
)

# Header Title
ws.merge_cells("A1:M1")
ws["A1"] = "유럽 전 법인 유통재고 주수(WOS) 리스크 진단 및 OLED 신모델 비중 비교 요약 (가명 처리 예제)"
ws["A1"].font = font_title
ws["A1"].alignment = Alignment(vertical="center")
ws.row_dimensions[1].height = 30

ws.merge_cells("A2:M2")
ws["A2"] = "기준 시점: 2026년 9월 (9/21W 주차 실적 반영, 9/28W 조회 기준) | 대외비 보안 준수를 위해 법인명 가명(A~P 법인) 처리"
ws["A2"].font = font_sub
ws.row_dimensions[2].height = 18

headers = [
    "구분", "법인코드", "법인명", "유통재고 WOS", "재고 리스크 판정",
    "유통재고(대)", "신모델 재고(대)", "Sell-out(대)", "OLED 신모델(대)",
    "OLED 내 신모델비중", "신모델 내 OLED비중", "전체TV 내 OLED신모델", "신모델 WOS", "구모델 WOS"
]

for col_num, h_text in enumerate(headers, 1):
    c = ws.cell(row=4, column=col_num, value=h_text)
    c.font = font_hdr
    c.fill = fill_hdr
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = thin_border
ws.row_dimensions[4].height = 26

row_idx = 5
for item in SUBS_MAPPING:
    is_eu = item.get("is_eu", False)
    is_risk = item["is_risk"]
    status_tag = "본부" if is_eu else ("★ 리스크" if is_risk else "정상")

    row_vals = [
        status_tag,
        item["code"],
        item["name"],
        item["wos"],
        item["risk"],
        item["inv_ttl"],
        item["inv_new"],
        item["so_ttl"],
        item["so_oled_new"],
        item["oled_in_oled"],
        item["oled_in_new"],
        item["oled_new_in_ttl"],
        item["wos_new"],
        item["wos_old"]
    ]

    for col_num, val in enumerate(row_vals, 1):
        c = ws.cell(row=row_idx, column=col_num, value=val)
        c.font = font_data
        c.border = thin_border

        if is_eu:
            c.fill = fill_total
            c.font = font_bold
        elif row_idx % 2 == 0:
            c.fill = fill_zebra

        # Format & Alignment
        if col_num in [1, 2, 5]:
            c.alignment = Alignment(horizontal="center", vertical="center")
            if col_num == 5:
                if is_risk and not is_eu:
                    c.fill = fill_risk
                    c.font = font_risk
                elif not is_eu:
                    c.fill = fill_safe
                    c.font = font_safe
        elif col_num == 3:
            c.alignment = Alignment(horizontal="left", vertical="center")
            if is_eu: c.font = font_bold
        elif col_num in [4, 13, 14]:
            c.alignment = Alignment(horizontal="right", vertical="center")
            c.number_format = '0.00"주"'
            if col_num == 4 and is_risk and not is_eu:
                c.font = font_risk
        elif col_num in [6, 7, 8, 9]:
            c.alignment = Alignment(horizontal="right", vertical="center")
            c.number_format = '#,##0"대"'
        elif col_num in [10, 11, 12]:
            c.alignment = Alignment(horizontal="right", vertical="center")
            c.number_format = '0.0%'

    ws.row_dimensions[row_idx].height = 20
    row_idx += 1

for col in ws.columns:
    max_len = 0
    col_letter = get_column_letter(col[0].column)
    for cell in col:
        v = str(cell.value or '')
        k_len = sum(2 if ord(ch) > 127 else 1 for ch in v)
        max_len = max(max_len, k_len)
    ws.column_dimensions[col_letter].width = max(max_len + 3, 11)

ws.column_dimensions['A'].width = 11
ws.column_dimensions['B'].width = 10
ws.column_dimensions['C'].width = 22
ws.column_dimensions['E'].width = 16

out_path = r'd:\TV 유럽영업\15. AX Task\2026 AX 실행과제\99. Master Agent\유럽_법인별_유통재고WOS_OLED신모델비중_요약.xlsx'
wb.save(out_path)
print(f"Masked Excel generated successfully at: {out_path}")
