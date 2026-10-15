import os
from pptx import Presentation

pptx_path = r"d:\TV 유럽영업\15. AX Task\2026 AX 실행과제\07. 제품 소개 사이트 자동 제작 에이전트\uploads\260619_ATA 가이드 v1.0_배포용.pptx"
prs = Presentation(pptx_path)

with open("ata_slides_text.txt", "w", encoding="utf-8") as out:
    for idx, slide in enumerate(prs.slides):
        out.write(f"\n================ Slide {idx+1} ================\n")
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    text = p.text.strip()
                    if text:
                        out.write(f"  [Text] {text}\n")
            if shape.has_table:
                out.write("  [Table]\n")
                table = shape.table
                for r_idx, row in enumerate(table.rows):
                    row_vals = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
                    out.write(f"    R{r_idx}: {' | '.join(row_vals)}\n")

print("Saved all slides to ata_slides_text.txt")
